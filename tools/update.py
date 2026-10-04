#!/usr/bin/env python3
"""Gera a dashboard a partir de home.xlsx (exportação da folha «Controlo Despesas HOME»).

Uso:  DASH_PW=<palavra-passe> python3 tools/update.py caminho/para/home.xlsx
Saídas:
  build/page.html    página em claro (para publicar no artifact do Claude; NÃO fazer commit)
  public/index.html  página cifrada (AES-256-GCM, PBKDF2-SHA256) servida pelo Vercel
"""
import os,sys,json,base64,collections,datetime,zoneinfo,warnings
import openpyxl
from cryptography.hazmat.primitives.ciphers.aead import AESGCM
from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC
from cryptography.hazmat.primitives import hashes
warnings.filterwarnings('ignore')
ROOT=os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
xlsx=sys.argv[1]; pw=os.environ['DASH_PW']
wb=openpyxl.load_workbook(xlsx,data_only=True)
def head(ws,n): return [str(c or '').strip() for c in next(ws.iter_rows(values_only=True))][:n]
assert head(wb['Obra'],12)==['Data','Entidade','Descrição','Valor','IVA','IS','IMT','TOTAL','Acumulado','Vencimento','Pago?','Evasão Fiscal'],'colunas da folha Obra mudaram'
assert head(wb['Projeto'],9)==['Data','Entidade','Descrição','Valor','IVA','IS','IMT','TOTAL','Acumulado'],'colunas da folha Projeto mudaram'
assert head(wb['Empréstimos'],6)==['Entidade','Saldo Inicial','Disponível','Consumido','Pago','Em Dívida'],'colunas da folha Empréstimos mudaram'
fmt=[c[0].number_format for c in wb['Projeto'].iter_rows()][1:]
R=[];skipped=[];chk=0
for r in list(wb['Obra'].iter_rows(values_only=True))[1:]:
    if r[0] is None: continue
    if r[3] is None:
        skipped.append(f"{r[0].strftime('%d/%m/%Y')}, {r[1]}, {r[2]}"); continue
    chk+=r[7] or 0
    R.append([r[0].strftime('%Y-%m-%d'),None,'Obra',r[1].strip(),r[2].strip(),round(r[3] or 0,2),round(r[4] or 0,2),round((r[5] or 0)+(r[6] or 0),2),round(r[7] or 0,2),r[9].strftime('%Y-%m-%d'),1 if r[10] else 0,round(r[11] or 0,2)])
for i,r in enumerate(list(wb['Projeto'].iter_rows(values_only=True))[1:]):
    if r[0] is None: continue
    chk+=r[7] or 0
    R.append([r[0].strftime('%Y-%m-%d'),r[0].strftime('%m/%Y') if 'd' not in fmt[i] else None,'Projeto',r[1].strip(),r[2].strip(),round(r[3] or 0,2),round(r[4] or 0,2),round((r[5] or 0)+(r[6] or 0),2),round(r[7] or 0,2),None,1,0])
cnt=collections.Counter(r[4] for r in R); best={}
for c,n in cnt.items():
    k=c.lower()
    if k not in best or n>cnt[best[k]]: best[k]=c
for r in R: r[4]=best[r[4].lower()]
emp=[[r[0]]+[round(x or 0,2) for x in r[1:6]] for r in list(wb['Empréstimos'].iter_rows(values_only=True))[1:] if r[0]]
tot=sum(r[8] for r in R)
assert len(R)>200,'menos de 200 lançamentos: exportação suspeita'
assert abs(tot-chk)<0.5,'soma dos totais não coincide com a folha'
asof=datetime.datetime.now(zoneinfo.ZoneInfo('Atlantic/Azores')).strftime('%Y-%m-%d')
data={'rows':R,'emp':emp,'skipped':skipped,'asof':asof}
tpl=open(os.path.join(ROOT,'tools','page.tpl.html'),encoding='utf-8').read()
assert tpl.count('/*DATA*/')==1
page=tpl.replace('/*DATA*/',json.dumps(data,ensure_ascii=False,separators=(',',':')))
os.makedirs(os.path.join(ROOT,'build'),exist_ok=True)
open(os.path.join(ROOT,'build','page.html'),'w',encoding='utf-8').write(page)
full='<!doctype html><html lang="pt-PT"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1,viewport-fit=cover"><meta name="robots" content="noindex,nofollow"><meta name="theme-color" content="#191c20"><meta name="apple-mobile-web-app-capable" content="yes"><meta name="apple-mobile-web-app-title" content="Despesas HOME"><style>html{color-scheme:dark}:root{padding-top:env(safe-area-inset-top,0px);padding-bottom:env(safe-area-inset-bottom,0px)}body{margin:0}img{max-width:100%}[hidden]{display:none!important}</style></head><body>'+page+'</body></html>'
IT=600000; salt,iv=os.urandom(16),os.urandom(12)
key=PBKDF2HMAC(hashes.SHA256(),32,salt,IT).derive(pw.encode())
b=lambda x:base64.b64encode(x).decode()
payload=json.dumps({'s':b(salt),'i':b(iv),'c':b(AESGCM(key).encrypt(iv,full.encode(),None)),'n':IT})
lock=open(os.path.join(ROOT,'tools','lock.html'),encoding='utf-8').read()
assert lock.count('/*PAYLOAD*/')==1
os.makedirs(os.path.join(ROOT,'public'),exist_ok=True)
open(os.path.join(ROOT,'public','index.html'),'w',encoding='utf-8').write(lock.replace('/*PAYLOAD*/',payload))
paid=sum(r[8] for r in R if r[10])
print(json.dumps({'asof':asof,'lancamentos':len(R),'investimento':round(tot,2),'pago':round(paid,2),'por_pagar':round(tot-paid,2),'emprestimos':emp,'sem_valor':skipped},ensure_ascii=False))
